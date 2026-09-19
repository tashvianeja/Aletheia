from __future__ import annotations

import io

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from privacy_guardian.analysis.documents import extract as extraction
from privacy_guardian.analysis.documents.extract import extract_document


def test_docx_extracts_paragraph_table_header_and_footer() -> None:
    document = Document()
    document.add_paragraph("Synthetic medical record")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Diagnosis"
    table.cell(0, 1).text = "Synthetic condition"
    section = document.sections[0]
    section.header.paragraphs[0].text = "Synthetic header"
    section.footer.paragraphs[0].text = "Synthetic footer"
    output = io.BytesIO()
    document.save(output)

    extracted = extract_document(output.getvalue(), "record.docx")

    assert extracted.partial is False
    assert extracted.document_type == "medical_record"
    assert len(extracted.pages) == 1
    assert all(
        text in extracted.pages[0].text
        for text in (
            "Synthetic medical record",
            "Diagnosis\tSynthetic condition",
            "Synthetic header",
            "Synthetic footer",
        )
    )


def test_xlsx_extracts_each_worksheet_and_marks_text_limit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "Identity"
    first.append(["Passport", "X00000001"])
    second = workbook.create_sheet("Finance")
    second.append(["Bank statement", "Synthetic"])
    output = io.BytesIO()
    workbook.save(output)
    monkeypatch.setattr(extraction, "MAX_TEXT", 18)

    extracted = extract_document(output.getvalue(), "records.xlsx")

    assert len(extracted.pages) == 2
    assert extracted.pages[0].number == 1
    assert extracted.pages[1].number == 2
    assert "Passport" in extracted.pages[0].text
    assert "Bank statement" in extracted.pages[1].text
    assert extracted.partial is True


def test_pptx_extracts_slide_text_and_table() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text_frame.text = "Synthetic resume work experience"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "Education"
    table.cell(0, 1).text = "Test University"
    output = io.BytesIO()
    presentation.save(output)

    extracted = extract_document(output.getvalue(), "resume.pptx")

    assert extracted.document_type == "resume"
    assert "Synthetic resume work experience" in extracted.pages[0].text
    assert "Education\tTest University" in extracted.pages[0].text


@pytest.mark.parametrize("suffix", ["txt", "csv", "json", "md", "xml", "py", "js"])
def test_text_formats_decode_non_utf8_completely_despite_expansion_limit(
    suffix: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(extraction, "MAX_TEXT", 20)
    text = "Résumé synthetic account number"
    content = text.encode("cp1252")

    extracted = extract_document(content, f"document.{suffix}")

    assert extracted.pages[0].text == text
    assert extracted.partial is False


def test_plain_text_above_byte_bound_samples_first_and_last_bytes(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(extraction, "MAX_BYTES", 20)
    monkeypatch.setattr(extraction, "SAMPLE_BYTES", 5)
    content = b"AAAAA01234567890ZZZZZ"

    extracted = extract_document(content, "oversized.txt")

    assert extracted.pages[0].text == "AAAAAZZZZZ"
    assert extracted.partial is True
    assert extracted.warnings == [
        "Large document sampled: first and last 5 MB; unsampled content was not checked."
    ]


def test_image_ocr_timeout_is_partial_without_leaking_parser_error(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    image = Image.new("RGB", (40, 20), "white")
    output = io.BytesIO()
    image.save(output, format="JPEG")

    def unavailable(*_args: object, **_kwargs: object) -> None:
        raise RuntimeError("synthetic secret from OCR backend")

    monkeypatch.setattr(extraction, "_ocr", unavailable)
    extracted = extract_document(output.getvalue(), "scan.jpg")

    assert extracted.has_images is True
    assert extracted.partial is True
    assert extracted.pages == []
    assert extracted.warnings == ["OCR unavailable or timed out; image text was not fully checked."]
    assert "synthetic secret" not in repr(extracted)


def test_original_size_and_unsupported_format_report_safe_partial_results() -> None:
    sampled = extract_document(
        b"synthetic text",
        "sample.txt",
        partial=True,
        original_size=extraction.MAX_BYTES + 1,
    )
    unsupported = extract_document(b"opaque synthetic bytes", "archive.unknown")

    assert sampled.partial is True
    assert "first and last 5 MB" in sampled.warnings[0]
    assert unsupported.partial is True
    assert unsupported.pages == []
    assert unsupported.warnings == [
        "Unsupported document format; no complete content scan was possible."
    ]


def test_archive_expansion_limit_fails_closed_without_parser_details(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        extraction.zipfile,
        "ZipFile",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(ValueError("synthetic archive internals")),
    )

    extracted = extract_document(b"PK synthetic", "damaged.docx")

    assert extracted.partial is True
    assert extracted.pages == []
    assert extracted.warnings == [
        "Document could not be completely extracted; it may be damaged, encrypted, or sampled."
    ]
    assert "synthetic archive internals" not in repr(extracted)
