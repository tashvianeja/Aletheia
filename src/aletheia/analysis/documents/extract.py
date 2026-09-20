from __future__ import annotations

import io
import re
import signal
import sys
import time
import zipfile
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from aletheia.core.events import DataCategory

MAX_BYTES = 50 * 1024 * 1024
SAMPLE_BYTES = 5 * 1024 * 1024
MAX_TEXT = 2_000_000
MAX_PAGES = 500


@dataclass
class Page:
    text: str
    number: int
    # OCR token boxes are worker-private and never serialised to the service.
    boxes: list[tuple[int, int, int, int, int, int]] = field(default_factory=list)


@dataclass
class ExtractedDocument:
    pages: list[Page] = field(default_factory=list)
    metadata_categories: set[DataCategory] = field(default_factory=set)
    partial: bool = False
    warnings: list[str] = field(default_factory=list)
    document_type: str = "generic"
    has_images: bool = False
    # Which national scheme an identity document belongs to, when it says so. What
    # may be covered on an ID and what has to stay readable is set by the scheme
    # that issued it, not by the fact that it is an ID: see documents/masking.py.
    id_scheme: str = ""


class AnalysisTimeout(TimeoutError):
    pass


@contextmanager
def _deadline(seconds: float) -> Iterator[None]:
    previous: Any = None
    armed = sys.platform != "win32"
    if sys.platform != "win32":
        try:
            previous = signal.getsignal(signal.SIGALRM)

            def handler(signum: int, frame: Any) -> None:
                raise AnalysisTimeout("Document analysis time budget reached")

            signal.signal(signal.SIGALRM, handler)
            signal.setitimer(signal.ITIMER_REAL, max(0.001, seconds))
        except ValueError:
            armed = False
    try:
        yield
    finally:
        if sys.platform != "win32" and armed:
            signal.setitimer(signal.ITIMER_REAL, 0)
            signal.signal(signal.SIGALRM, previous)


def classify_document(text: str, suffix: str = "") -> str:
    lower = text.lower()
    choices = {
        "identity_document": (
            r"passport",
            r"national identity",
            r"P<[A-Z<]{3}",
            r"driver.s licen[cs]e",
            r"aadhaar",
            r"आधार",
            r"unique identification authority",
        ),
        "bank_statement": (
            r"bank statement",
            r"account number",
            r"opening balance",
            r"closing balance",
        ),
        "medical_record": (r"diagnosis", r"prescribed", r"patient name", r"medical record"),
        "resume": (r"curriculum vitae", r"work experience", r"education", r"resume"),
        "tax_form": (r"tax return", r"form w-?2", r"form 1040", r"taxpayer"),
        "source_code": (r"\bdef \w+\(", r"\bimport \w+", r"\bfunction\s+\w+\(", r"\bclass \w+"),
    }
    scores = [
        (
            sum(
                bool(re.search(pattern.lower(), lower))
                if any(character in pattern for character in r"\.^$*+?{}[]|()")
                else pattern.lower() in lower
                for pattern in patterns
            ),
            category,
        )
        for category, patterns in choices.items()
    ]
    best = max(scores)
    if best[0]:
        return best[1]
    return (
        "photo"
        if suffix in {".png", ".jpg", ".jpeg", ".webp", ".tiff", ".bmp", ".heic"}
        else "generic"
    )


def identity_scheme(text: str) -> str:
    """The scheme that issued an identity document, where the document names it.

    Only Aadhaar is recognised by name so far. Every other ID falls back to "", and
    the redaction rules for an unrecognised scheme are the cautious ones.
    """
    if re.search(r"aadhaar|aadhar|आधार|uidai|unique\s+identification\s+authority", text, re.I):
        return "aadhaar"
    return ""


def _ocr(image: Any, number: int, remaining: float) -> Page:
    import pytesseract

    data = pytesseract.image_to_data(
        image.convert("RGB"),
        output_type=pytesseract.Output.DICT,
        timeout=max(0.1, remaining),
        # Automatic page segmentation, which is what a page is. Reading a page as one
        # uniform block instead drops whatever does not share the body's layout, and on
        # an identity card the line it drops is the number printed large across the
        # middle: the one detail on the card that most needs a box over it.
        config="--psm 3",
    )
    text = ""
    boxes = []
    previous_line: tuple[int, int, int] | None = None
    for index, word in enumerate(data["text"]):
        word = str(word).strip()
        if not word:
            continue
        line = (
            int(data["block_num"][index]),
            int(data["par_num"][index]),
            int(data["line_num"][index]),
        )
        if text:
            text += "\n" if line != previous_line else " "
        start = len(text)
        text += word
        boxes.append(
            (
                start,
                len(text),
                int(data["left"][index]),
                int(data["top"][index]),
                int(data["width"][index]),
                int(data["height"][index]),
            )
        )
        previous_line = line
    return Page(text, number, boxes)


def _check_archive(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        infos = archive.infolist()
        if len(infos) > 20_000 or sum(item.file_size for item in infos) > 150 * 1024 * 1024:
            raise ValueError("Archive expansion limit reached")
        if any(
            item.file_size > 10 * 1024 * 1024 and item.file_size / max(1, item.compress_size) > 300
            for item in infos
        ):
            raise ValueError("Archive expansion ratio exceeds safe limit")


def extract_document(
    data: bytes,
    filename: str = "document.txt",
    *,
    timeout: float = 20,
    partial: bool = False,
    original_size: int | None = None,
) -> ExtractedDocument:
    result = ExtractedDocument(partial=partial)
    suffix = Path(filename).suffix.lower()
    if len(data) > MAX_BYTES:
        data = data[:SAMPLE_BYTES] + data[-SAMPLE_BYTES:]
        result.partial = True
    if result.partial or (original_size is not None and original_size > MAX_BYTES):
        result.partial = True
        result.warnings.append(
            "Large document sampled: first and last 5 MB; unsampled content was not checked."
        )
    start = time.monotonic()

    def remaining() -> float:
        value = timeout - (time.monotonic() - start)
        if value <= 0:
            raise AnalysisTimeout()
        return value

    try:
        with _deadline(timeout):
            if suffix == ".pdf" or data.startswith(b"%PDF-"):
                from pypdf import PdfReader

                reader = PdfReader(io.BytesIO(data), strict=False)
                if reader.is_encrypted and not reader.decrypt(""):
                    raise ValueError("Encrypted document requires a password")
                for index, pdf_page in enumerate(reader.pages):
                    remaining()
                    if index >= MAX_PAGES:
                        result.partial = True
                        result.warnings.append("Page limit reached.")
                        break
                    result.has_images = result.has_images or bool(pdf_page.images)
                    text = pdf_page.extract_text() or ""
                    if (
                        classify_document(text) in {"bank_statement", "tax_form"}
                        and remaining() > 1
                    ):
                        import pdfplumber

                        with pdfplumber.open(io.BytesIO(data)) as layout:
                            tables = layout.pages[index].extract_tables()
                            table_text = "\n".join(
                                "\t".join(str(cell or "") for cell in row)
                                for table in tables
                                for row in table
                            )
                            if table_text:
                                text += "\n" + table_text
                    if len(text) > MAX_TEXT:
                        result.partial = True
                    if text.strip():
                        result.pages.append(Page(text[:MAX_TEXT], index + 1))
                    else:
                        import pypdfium2

                        try:
                            with pypdfium2.PdfDocument(data) as raster:
                                image = raster[index].render(scale=2).to_pil()
                                result.pages.append(_ocr(image, index + 1, remaining()))
                        except (RuntimeError, OSError, ValueError) as error:
                            # Without OCR a scanned page must be reported as unchecked, never
                            # allowed to fail the whole analysis and let the file through.
                            result.partial = True
                            if not result.warnings:
                                result.warnings.append(
                                    "OCR unavailable or timed out; "
                                    "scanned pages were not fully checked."
                                )
                            del error
                    if sum(len(page.text) for page in result.pages) >= MAX_TEXT:
                        result.partial = True
                        break
            elif suffix == ".docx":
                _check_archive(data)
                from docx import Document

                document = Document(io.BytesIO(data))
                lines = [paragraph.text for paragraph in document.paragraphs]
                for table in document.tables:
                    lines.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
                for section in document.sections:
                    lines.extend(paragraph.text for paragraph in section.header.paragraphs)
                    lines.extend(paragraph.text for paragraph in section.footer.paragraphs)
                text = "\n".join(lines)
                result.partial = result.partial or len(text) > MAX_TEXT
                result.pages.append(Page(text[:MAX_TEXT], 1))
            elif suffix in {".xlsx", ".xlsm"}:
                _check_archive(data)
                from openpyxl import load_workbook

                workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
                try:
                    for index, sheet in enumerate(workbook.worksheets):
                        lines = []
                        size = 0
                        for row in sheet.iter_rows(values_only=True):
                            remaining()
                            line = "\t".join(
                                str(value) if value is not None else "" for value in row
                            )
                            lines.append(line)
                            size += len(line)
                            if size >= MAX_TEXT:
                                result.partial = True
                                break
                        text = "\n".join(lines)
                        result.partial = result.partial or len(text) > MAX_TEXT
                        result.pages.append(Page(text[:MAX_TEXT], index + 1))
                finally:
                    workbook.close()
            elif suffix == ".pptx":
                _check_archive(data)
                from pptx import Presentation

                presentation = Presentation(io.BytesIO(data))
                for index, slide in enumerate(presentation.slides):
                    remaining()
                    lines = []
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            lines.append(shape.text_frame.text)
                        if shape.has_table:
                            lines.extend(
                                "\t".join(cell.text for cell in row.cells)
                                for row in shape.table.rows
                            )
                    if slide.has_notes_slide:
                        lines.append(slide.notes_slide.notes_text_frame.text)
                    text = "\n".join(lines)
                    result.partial = result.partial or len(text) > MAX_TEXT
                    result.pages.append(Page(text[:MAX_TEXT], index + 1))
            elif suffix in {".png", ".jpg", ".jpeg", ".webp", ".tiff", ".tif", ".bmp"}:
                from PIL import Image

                result.has_images = True
                with Image.open(io.BytesIO(data)) as image:
                    if image.width * image.height > 40_000_000:
                        raise ValueError("Image pixel limit reached")
                    exif = image.getexif()
                    if exif.get(34853):
                        result.metadata_categories.add(DataCategory.LOCATION_PRECISE)
                    if any(exif.get(tag) for tag in (315, 40093)):
                        result.metadata_categories.add(DataCategory.FULL_NAME)
                    try:
                        result.pages.append(_ocr(image, 1, remaining()))
                    except (RuntimeError, OSError) as error:
                        result.partial = True
                        result.warnings.append(
                            "OCR unavailable or timed out; image text was not fully checked."
                        )
                        del error
            elif suffix in {
                ".txt",
                ".csv",
                ".json",
                ".md",
                ".log",
                ".xml",
                ".html",
                ".htm",
                ".yaml",
                ".yml",
                ".py",
                ".js",
                ".ts",
                ".ini",
                ".rtf",
                "",
            }:
                import chardet

                encoding = str(chardet.detect(data[:100_000]).get("encoding") or "utf-8")
                text = data.decode(encoding, errors="replace")
                # The byte limit already bounds ordinary text. Expanded archive/PDF
                # content retains the stricter MAX_TEXT guard in its own branches.
                result.pages.append(Page(text, 1))
            else:
                result.partial = True
                result.warnings.append(
                    "Unsupported document format; no complete content scan was possible."
                )
    except AnalysisTimeout:
        result.partial = True
        result.warnings.append("Analysis time limit reached; results cover only completed pages.")
    except Exception:
        # Third-party parser exception messages may quote source contents; never propagate them.
        result.partial = True
        result.warnings.append(
            "Document could not be completely extracted; it may be damaged, encrypted, or sampled."
        )
        if not result.pages and result.partial and suffix == ".pdf":
            fragments = re.findall(
                rb"[\x20-\x7e]{8,200}", data[: 5 * 1024 * 1024] + data[-5 * 1024 * 1024 :]
            )
            result.pages.append(
                Page("\n".join(fragment.decode("ascii") for fragment in fragments)[:MAX_TEXT], 1)
            )
    if result.partial and not result.warnings:
        result.warnings.append("Extraction limit reached; omitted content has not been checked.")
    all_text = "\n".join(page.text for page in result.pages)
    result.document_type = classify_document(all_text, suffix)
    if result.document_type == "identity_document":
        result.id_scheme = identity_scheme(all_text)
        if result.has_images:
            result.metadata_categories.add(DataCategory.BIOMETRIC_PHOTO)
    return result
